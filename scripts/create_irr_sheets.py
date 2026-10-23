"""
Generate IRR labeling spreadsheets for inter-rater reliability.

Pulls a 10% stratified sample from the 30% LLM-labeled set and creates
two independent Excel workbooks (one per annotator). Each sheet has
the review text, a dropdown for the human label, and hidden columns
with the LLM's label and confidence (to compare after annotation).

Run: python scripts/create_irr_sheets.py
"""

# Windows consoles default to cp1252; bot names and review text are not.
import sys as _sys

for _stream in (_sys.stdout, _sys.stderr):
    if hasattr(_stream, "reconfigure"):
        try:
            _stream.reconfigure(encoding="utf-8", errors="replace")
        except (ValueError, OSError):
            pass

import pandas as pd
from sklearn.model_selection import train_test_split
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

df_full = pd.read_csv(
    "data/labeled/experiment_labels_accumulated.csv",
    low_memory=False,
    encoding="latin-1"
)

# First 54,942 rows = the 30% batch (labels were added incrementally)
df_full = df_full.iloc[:54942].copy()
df_full = df_full.rename(columns={"llm_label": "label", "llm_confidence": "confidence"})
df_full = df_full[df_full["label"].isin(["privacy", "not_privacy"])].reset_index(drop=True)

# Stratified 10% sample
_, df = train_test_split(
    df_full,
    test_size=0.10,
    stratify=df_full["label"],
    random_state=42
)
df = df.reset_index(drop=True)

print(f"Full 30% set:    {len(df_full):,} reviews")
print(f"IRR 10% sample:  {len(df):,} reviews")
print(f"  Privacy:       {(df['label'] == 'privacy').sum():,}")
print(f"  Not-privacy:   {(df['label'] == 'not_privacy').sum():,}")

header_fill = PatternFill("solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", name="Arial", size=11)
border = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)


def create_file(filename, annotator_name):
    print(f"\nCreating {filename}...")
    wb = Workbook()

    ws_inst = wb.active
    ws_inst.title = "Instructions"
    ws_inst.column_dimensions["A"].width = 90

    instructions = [
        (f"INTER-RATER RELIABILITY (IRR) — {annotator_name.upper()}", True, 14),
        ("", False, 11),
        (f"Annotator: {annotator_name}", False, 11),
        (f"Total reviews: {len(df):,}  |  Privacy: {(df['label']=='privacy').sum():,}  |  Not-privacy: {(df['label']=='not_privacy').sum():,}", False, 11),
        ("Source: 10% stratified sample of 30%-labeled set", False, 11),
        ("", False, 11),
        ("STEPS", True, 12),
        ("1. Go to the 'Labeling' sheet.", False, 11),
        ("2. Read each review in column B.", False, 11),
        ("3. Select your label from the dropdown in column C: privacy or not_privacy", False, 11),
        ("4. Columns D and E are hidden (LLM label & confidence). Do NOT unhide until done.", False, 11),
        ("5. After finishing, unhide D-E to compare with LLM labels.", False, 11),
        ("6. Return completed file to lead annotator for kappa computation.", False, 11),
        ("", False, 11),
        ("LABEL DEFINITIONS", True, 12),
        ("  privacy     = Review mentions ANY privacy topic:", False, 11),
        ("                data collection, sharing, permissions, tracking, DMs,", False, 11),
        ("                authentication, surveillance, consent, data retention, etc.", False, 11),
        ("  not_privacy = Review discusses only: functionality, bugs, UI,", False, 11),
        ("                performance, pricing, or general praise/complaints.", False, 11),
        ("", False, 11),
        ("IMPORTANT: Label independently. Do not discuss with the other annotator.", True, 11),
    ]

    for i, (line, bold, size) in enumerate(instructions, 1):
        cell = ws_inst.cell(row=i, column=1, value=line)
        cell.font = Font(bold=bold, size=size, name="Arial")
        cell.alignment = Alignment(wrap_text=True)

    ws = wb.create_sheet("Labeling")
    headers = ["#", "Review Text", f"{annotator_name} Label", "LLM Label", "LLM Confidence"]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    ws.row_dimensions[1].height = 22
    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 85
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16

    for i, row in df.iterrows():
        r = i + 2
        row_fill = PatternFill("solid", fgColor="F7FAFF") if i % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")

        id_cell = ws.cell(row=r, column=1, value=i + 1)
        id_cell.border = border
        id_cell.font = Font(name="Arial", size=10)
        id_cell.alignment = Alignment(horizontal="center", vertical="top")
        id_cell.fill = row_fill

        text = str(row["text"])[:500] if pd.notna(row["text"]) else ""
        text_cell = ws.cell(row=r, column=2, value=text)
        text_cell.border = border
        text_cell.alignment = Alignment(wrap_text=True, vertical="top")
        text_cell.font = Font(name="Arial", size=10)
        text_cell.fill = row_fill

        label_cell = ws.cell(row=r, column=3, value="")
        label_cell.border = border
        label_cell.font = Font(name="Arial", size=10)
        label_cell.alignment = Alignment(horizontal="center", vertical="top")
        label_cell.fill = PatternFill("solid", fgColor="EBF3FB")

        llm_cell = ws.cell(row=r, column=4, value=row["label"])
        llm_cell.border = border
        llm_cell.font = Font(name="Arial", size=10, color="888888")
        llm_cell.alignment = Alignment(horizontal="center", vertical="top")

        conf = round(row["confidence"], 3) if "confidence" in row and pd.notna(row["confidence"]) else ""
        conf_cell = ws.cell(row=r, column=5, value=conf)
        conf_cell.border = border
        conf_cell.font = Font(name="Arial", size=10, color="888888")
        conf_cell.alignment = Alignment(horizontal="center", vertical="top")

        if i % 1000 == 0:
            print(f"  Written {i:,} rows...")

    dv = DataValidation(type="list", formula1='"privacy,not_privacy"', allow_blank=True)
    dv.prompt = "Select label"
    dv.promptTitle = "Label"
    ws.add_data_validation(dv)
    dv.add(f"C2:C{len(df) + 1}")

    ws.column_dimensions["D"].hidden = True
    ws.column_dimensions["E"].hidden = True
    ws.freeze_panes = "B2"

    ws_sum = wb.create_sheet("Summary")
    ws_sum.column_dimensions["A"].width = 35
    ws_sum.column_dimensions["B"].width = 20

    last_row = len(df) + 1
    summary_data = [
        ("IRR Summary", None),
        ("", None),
        ("Annotator", annotator_name),
        ("Total reviews (10% of 30% set)", len(df)),
        ("LLM: Privacy", int((df["label"] == "privacy").sum())),
        ("LLM: Not-privacy", int((df["label"] == "not_privacy").sum())),
        ("LLM Privacy Rate (%)", f"{(df['label']=='privacy').mean()*100:.2f}%"),
        ("Avg LLM Confidence", f"{df['confidence'].mean():.3f}" if "confidence" in df.columns else "N/A"),
        ("", None),
        ("Human labels completed", f"=COUNTA(Labeling!C2:C{last_row})"),
        ("Human: Privacy count", f'=COUNTIF(Labeling!C2:C{last_row},"privacy")'),
        ("Human: Not-privacy count", f'=COUNTIF(Labeling!C2:C{last_row},"not_privacy")'),
        ("% Complete", f"=COUNTA(Labeling!C2:C{last_row})/{len(df)}"),
    ]

    for i, (label, val) in enumerate(summary_data, 1):
        lc = ws_sum.cell(row=i, column=1, value=label)
        lc.font = Font(name="Arial", size=11, bold=(i == 1))
        if val is not None:
            vc = ws_sum.cell(row=i, column=2, value=val)
            vc.font = Font(name="Arial", size=11)
            vc.alignment = Alignment(horizontal="right")

    wb.save(filename)
    print(f"  Saved: {filename}")


create_file("IRR_Annotator1.xlsx", "Annotator 1")
create_file("IRR_Annotator2.xlsx", "Annotator 2")

print(f"\nDone! Two files created with {len(df):,} reviews each.")
print("After both annotators finish, run: python compute_kappa.py")